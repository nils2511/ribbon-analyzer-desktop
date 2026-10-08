"""Security regressions use fabricated data and never contact remote APIs."""
import csv
import json
import tempfile
import unittest
from pathlib import Path
from openpyxl import load_workbook
from src.models import Project, ImageResult, RibbonAnnotation
from src.export_module import export_csv, export_excel


def project_with_note(note, age='P7', genotype='WT'):
    image = ImageResult('sample.tif', 'sample.tif', 'sample', age, genotype, analyzed=True)
    image.ribbons = [RibbonAnnotation(1, 'sample.tif', 'sample.tif', 'sample', age,
                                    genotype, notes=note, length_nm=200)]
    return Project('synthetic', '.', image_results=[image])


class ExportSecurityTests(unittest.TestCase):
    def test_excel_preserves_formula_like_notes_as_literal_text(self):
        for note in ['=1+1', '+1+1', '-1+1', '@SUM(1,1)']:
            with self.subTest(note=note), tempfile.TemporaryDirectory() as d:
                project = project_with_note(note)
                path = export_excel(project, Path(d))
                book = load_workbook(path, data_only=False)
                cells = list(book['All Ribbons'].iter_rows())
                index = next(i for i,c in enumerate(cells[0]) if c.value=='notes')
                self.assertEqual(cells[1][index].value, note)
                self.assertNotEqual(cells[1][index].data_type, 'f')
                self.assertEqual(project.image_results[0].ribbons[0].notes, note)
                book.close()

    def test_csv_escapes_formula_and_control_prefixes(self):
        for note in ['=1+1', '+1+1', '-1+1', '@SUM(1,1)', '\t=1+1', '\r=1+1', '  =1+1']:
            with self.subTest(note=note), tempfile.TemporaryDirectory() as d:
                project = project_with_note(note)
                path = export_csv(project, Path(d))
                with path.open(newline='') as f:
                    row = next(csv.DictReader(f))
                self.assertTrue(row['notes'].startswith("'"))
                self.assertEqual(row['length_nm'], '200')
                self.assertEqual(project.image_results[0].ribbons[0].notes, note)

    def test_excel_handles_invalid_and_colliding_group_sheet_names(self):
        with tempfile.TemporaryDirectory() as d:
            project = project_with_note('plain', 'P7/early', 'WT:group')
            other = project_with_note('plain', 'P7_early', 'WT_group')
            project.image_results.extend(other.image_results)
            book = load_workbook(export_excel(project, Path(d)))
            self.assertEqual(len(book.sheetnames), 7)
            self.assertEqual(len(set(name.casefold() for name in book.sheetnames)), 7)
            book.close()


class ProjectSecurityTests(unittest.TestCase):
    def test_invalid_calibration_is_rejected_before_ui_use(self):
        for value in [0, -1, 'invalid', None, float('inf')]:
            with self.subTest(value=value), tempfile.TemporaryDirectory() as d:
                project = project_with_note('plain')
                data = json.loads(json.dumps(project, default=lambda value: value.__dict__))
                data['image_results'][0]['scale_px'] = value
                path = Path(d) / 'project.json'
                path.write_text(json.dumps(data))
                with self.assertRaises(ValueError):
                    Project.load(path)

    def test_malformed_geometry_is_rejected_before_ui_use(self):
        with tempfile.TemporaryDirectory() as d:
            project = project_with_note('plain')
            project.image_results[0].ribbons[0].polyline_pts = [[1], [2, 3]]
            path = Path(d) / 'project.json'
            project.save(path)
            with self.assertRaises(ValueError):
                Project.load(path)

    def test_relocation_cannot_escape_selected_root(self):
        with tempfile.TemporaryDirectory() as d:
            root = Path(d) / 'images'; root.mkdir()
            outside = Path(d) / 'private.tif'; outside.write_bytes(b'synthetic')
            project = project_with_note('plain')
            ir = project.image_results[0]
            project.root_folder = '/missing'
            ir.image_path = '/missing/profile.tif'
            ir.image_filename = '../private.tif'
            self.assertEqual(project.relocate_images(root), (0, 1))
            self.assertEqual(outside.read_bytes(), b'synthetic')


if __name__ == '__main__':
    unittest.main()
